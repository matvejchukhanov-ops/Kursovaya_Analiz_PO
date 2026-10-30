# Сборка текста курсовой работы в .docx (python-docx)
import os, re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SCR = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(SCR, "fig")
ASM = os.path.join(SCR, "user.asm")     # листинг Capstone пользовательского кода Whiter.a
PROJ = r"C:\Claude pr\kursovaya"
REPORT = os.path.join(PROJ, "build", "report_Whiter.a.exe.txt")
OUT = sys.argv[1]

doc = Document()

# ================================================================ стили

def set_fonts(rpr_owner, name, size=None, bold=None, italic=None, color=True):
    f = rpr_owner.font
    f.name = name
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if color: f.color.rgb = RGBColor(0, 0, 0)
    rpr = rpr_owner.element.get_or_add_rPr() if hasattr(rpr_owner, "element") else None
    if rpr is not None:
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts"); rpr.insert(0, rfonts)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(a)) is not None:
                del rfonts.attrib[qn(a)]
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(a), name)

st = doc.styles
normal = st["Normal"]
set_fonts(normal, "Times New Roman", 14)
pf = normal.paragraph_format
pf.line_spacing = 1.5; pf.first_line_indent = Cm(1.25); pf.space_after = Pt(0); pf.space_before = Pt(0)
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for name, size, align, indent in (("Heading 1", 16, WD_ALIGN_PARAGRAPH.CENTER, 0),
                                  ("Heading 2", 14, WD_ALIGN_PARAGRAPH.LEFT, 1.25),
                                  ("Heading 3", 14, WD_ALIGN_PARAGRAPH.LEFT, 1.25)):
    s = st[name]
    set_fonts(s, "Times New Roman", size, bold=True, italic=False)
    p = s.paragraph_format
    p.alignment = align; p.first_line_indent = Cm(indent); p.left_indent = Cm(0)
    p.space_before = Pt(12 if name != "Heading 1" else 0); p.space_after = Pt(12 if name == "Heading 1" else 6)
    p.line_spacing = 1.5; p.keep_with_next = True
    p.page_break_before = name == "Heading 1"

# оглавление: Word при обновлении поля TOC берёт оформление из стилей «toc N»
for lvl, ind in ((1, 0.0), (2, 0.5), (3, 1.0)):
    s = st.add_style(f"toc {lvl}", 1)
    s.base_style = normal
    p = s.paragraph_format
    p.first_line_indent = Cm(0); p.left_indent = Cm(ind); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.line_spacing = 1.5; p.space_after = Pt(0)

code = st.add_style("Code", 1)
set_fonts(code, "Courier New", 9)
# заливку — до остальных свойств абзаца: порядок элементов pPr задан схемой
shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "F2F2F2")
code.element.get_or_add_pPr().append(shd)
cp = code.paragraph_format
cp.line_spacing = 1.0; cp.first_line_indent = Cm(0); cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
cp.space_after = Pt(0); cp.space_before = Pt(0)
code_small = st.add_style("CodeSmall", 1)
code_small.base_style = code
set_fonts(code_small, "Courier New", 8)

sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)
sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
sec.different_first_page_header_footer = True

# номер страницы внизу по центру
def add_field(paragraph, instr):
    r = paragraph.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); r._r.append(b)
    r2 = paragraph.add_run()
    t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = instr; r2._r.append(t)
    r3 = paragraph.add_run()
    s_ = OxmlElement("w:fldChar"); s_.set(qn("w:fldCharType"), "separate"); r3._r.append(s_)
    r4 = paragraph.add_run("1")
    r5 = paragraph.add_run()
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r5._r.append(e)
    for r_ in (r, r2, r3, r4, r5):
        r_.font.name = "Times New Roman"; r_.font.size = Pt(12)

fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.first_line_indent = Cm(0)
add_field(fp, "PAGE")

# ================================================================ помощники

fig_no = [0]; tab_no = [0]; lst_no = [0]


def P(text, align=None, indent=True, bold=False, italic=False, size=None, space_after=None, keep=False):
    p = doc.add_paragraph()
    if align is not None: p.alignment = align
    if not indent: p.paragraph_format.first_line_indent = Cm(0)
    if space_after is not None: p.paragraph_format.space_after = Pt(space_after)
    if keep: p.paragraph_format.keep_with_next = True
    # **жирный** внутри текста
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part: continue
        r = p.add_run(part[2:-2] if part.startswith("**") else part)
        r.bold = bold or part.startswith("**")
        r.italic = italic
        if size: r.font.size = Pt(size)
    return p


def L(items, numbered=False):
    """Перечисление по ГОСТ: тире или «1)», с отступом первой строки."""
    for i, it in enumerate(items):
        end = ";" if i < len(items) - 1 else "."
        text = it.rstrip(".;")
        mark = f"{i + 1}) " if numbered else "– "
        P(mark + text + end)


def H1(text): return doc.add_heading(text, level=1)
def H2(text): return doc.add_heading(text, level=2)
def H3(text): return doc.add_heading(text, level=3)


def FIGURE(name, caption, width=16):
    fig_no[0] += 1
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(os.path.join(FIG, name), width=Cm(width))
    c = P(f"Рисунок {fig_no[0]} – {caption}", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_after=6)
    return fig_no[0]


def next_fig(): return fig_no[0] + 1
def next_tab(): return tab_no[0] + 1
def next_lst(): return lst_no[0] + 1


def set_cell_borders(table, none=False):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil" if none else "single")
        if not none:
            el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "000000")
        borders.append(el)
    successor = None
    for tag in ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook"):
        successor = tblPr.find(qn(tag))
        if successor is not None:
            break
    if successor is not None:
        successor.addprevious(borders)
    else:
        tblPr.append(borders)


def TABLE(caption, headers, rows, widths, size=12):
    tab_no[0] += 1
    cap = P(f"Таблица {tab_no[0]} – {caption}", align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, keep=True)
    cap.paragraph_format.space_before = Pt(6)
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_borders(t)
    for r, row in enumerate([headers] + rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.width = Cm(widths[c])
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if r == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(val))
            run.font.size = Pt(size)
            run.bold = r == 0
            if r == 0:
                tr = t.rows[0]._tr
                trPr = tr.get_or_add_trPr()
                h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return tab_no[0]


def LISTING(caption, lines, style="Code"):
    lst_no[0] += 1
    cap = P(f"Листинг {lst_no[0]} – {caption}", align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, keep=True)
    cap.paragraph_format.space_before = Pt(6)
    for i, line in enumerate(lines):
        p = doc.add_paragraph(style=style)
        if i < len(lines) - 1 and len(lines) < 45:
            p.paragraph_format.keep_with_next = True
        p.add_run(line.replace("\t", "    ") if line else " ")
    doc.add_paragraph()
    return lst_no[0]


def PAGEBREAK():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------- исходные данные

asm = {}
order = []
for line in open(ASM, encoding="utf-8"):
    m = re.match(r"([0-9A-F]{8})  (\S+)\s*(.*)", line.rstrip("\n"))
    if m:
        a = int(m.group(1), 16)
        asm[a] = (m.group(2), m.group(3))
        order.append(a)


def asm_lines(start, end, notes):
    out = []
    for a in order:
        if a < start or a > end:
            continue
        mn, ops = asm[a]
        if mn == "nop":
            continue
        ops = re.sub(r"0x([0-9a-f]+)", lambda m: m.group(1).upper() + "h", ops)
        text = f"{a:08X}  {mn:<6} {ops}"
        if a in notes:
            text = f"{text:<52}; {notes[a]}" if ";" not in text else f"{text}  — {notes[a]}"
        out.append(text)
    return out


def src(path, start_marker, end_marker, include_end=True):
    lines = open(path, encoding="utf-8").read().splitlines()
    s = next(i for i, l in enumerate(lines) if start_marker in l)
    e = next(i for i, l in enumerate(lines[s + 1:], s + 1) if end_marker in l)
    return lines[s:e + 1 if include_end else e]


report = open(REPORT, encoding="utf-8-sig").read().splitlines()


def report_block(start_marker, end_marker):
    s = next(i for i, l in enumerate(report) if start_marker in l)
    e = next(i for i, l in enumerate(report[s + 1:], s + 1) if end_marker in l)
    return report[s:e]


# ================================================================ титульный лист

def center(text, bold=False, size=14, space_after=0, spacing=1.15):
    p = P(text, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, bold=bold, size=size, space_after=space_after)
    p.paragraph_format.line_spacing = spacing
    return p

for line in ("МИНОБРНАУКИ РОССИИ",
             "Федеральное государственное бюджетное образовательное учреждение высшего образования",
             "«МИРЭА – Российский технологический университет»", "РТУ МИРЭА",
             "Институт кибербезопасности и цифровых технологий",
             "Кафедра КБ-2 «Информационно-аналитические системы кибербезопасности»"):
    center(line)
for _ in range(4): center("")
center("ОТЧЁТ ПО КУРСОВОЙ РАБОТЕ", bold=True, size=16)
center("по дисциплине «Анализ программного обеспечения на наличие вредоносных объектов»")
center("")
center("на тему:")
center("«Анализ взаимодействия вредоносной программы с файловой системой»", bold=True)
for _ in range(3): center("")

t = doc.add_table(rows=2, cols=2)
set_cell_borders(t, none=True)
cells = [("Отчёт представлен к рассмотрению:\nСтудент группы БББО-24-24",
          "«___» ____________ 2026 г.\n\n________________ (подпись)\nЧуханов М.Г."),
         ("Преподаватель",
          "«___» ____________ 2026 г.\n\n________________ (подпись)\n____________________")]
for r, (left, right) in enumerate(cells):
    for c, text in enumerate((left, right)):
        cell = t.cell(r, c); cell.width = Cm(8.2)
        cell.text = ""
        for k, line in enumerate(text.split("\n")):
            p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            p.add_run(line)
        cell.paragraphs[-1].paragraph_format.space_after = Pt(18)
for _ in range(5): center("")
center("Москва 2026")

# ================================================================ содержание

PAGEBREAK()
center("СОДЕРЖАНИЕ", bold=True, size=16, space_after=12, spacing=1.5)
toc = doc.add_paragraph(); toc.paragraph_format.first_line_indent = Cm(0)
add_field(toc, 'TOC \\o "1-3" \\h \\z \\u')

# ================================================================ введение

H1("ВВЕДЕНИЕ")
P("Файловая система — главный объект воздействия большинства вредоносных программ. Через неё вредоносное "
  "программное обеспечение (ВПО) закрепляется в системе, маскируется под системные компоненты, ищет ценные "
  "данные и, в наиболее опасном случае, уничтожает или шифрует их. Деструктивные программы-вайперы, "
  "которые безвозвратно стирают пользовательские файлы, применялись в крупнейших инцидентах последних лет "
  "(NotPetya в 2017 году, WhisperGate и HermeticWiper в 2022 году). Поэтому умение по исполняемому файлу "
  "понять, какие операции он выполняет над файлами и каталогами, — базовый навык специалиста по анализу ВПО.")
P("В цикле практических работ по дисциплине был создан инструмент статического анализа исполняемых файлов: "
  "извлечение строк (ПР №3), сопоставление сигнатур с классификаторами CWE и CVE (ПР №6), автоматическая "
  "«Карта инцидентов ИБ» IncidentMap (ПР №7). Практическая работа №8 ставит задачу развить его в прототип "
  "информационно-аналитической системы (ИАС) безопасности, которая по пути к файлу формирует структурированный "
  "отчёт. Курсовая работа продолжает этот цикл: ИАС дополнена модулем анализа взаимодействия программы с файловой "
  "системой, который работает непосредственно с машинным кодом и восстанавливает аргументы вызовов функций "
  "Windows API.")
P("**Цель работы** — исследовать механизмы взаимодействия вредоносной программы с файловой системой на примере "
  "образца Whiter.a и разработать средство, автоматически выявляющее такие механизмы методами статического анализа.")
P("Для достижения цели решаются задачи:", keep=True)
L(["рассмотреть способы взаимодействия ВПО с файловой системой Windows и их классификацию по матрице MITRE ATT&CK",
   "доработать средство IncidentMap из ПР №7 до ИАС по техническому заданию ПР №8 и добавить модуль анализа "
   "операций с файловой системой по машинному коду",
   "выполнить статический анализ образца Whiter.a: структура PE-файла, таблица импорта, строки",
   "восстановить алгоритм работы образца и последствия его запуска для файловой системы",
   "сформировать индикаторы компрометации и рекомендации по обнаружению, удалению и защите",
   "проверить ИАС на легитимных программах и оценить долю ложных срабатываний"], numbered=True)
P("**Объект исследования** — исполняемый файл Whiter.a.exe (57 344 байта, SHA-256 "
  "fe974c995cfb27e8c91123081986847f6d3d4252b6a8d1e1385c558f2aeb7057). **Предмет исследования** — операции "
  "образца над файлами, каталогами, дисками и связанными с ними ключами реестра.")
P("Образец относится к деструктивному ВПО, поэтому в работе применён только статический анализ: файл ни разу "
  "не запускался и даже не распаковывался из архива на диск — ИАС получает его содержимое через стандартный ввод. "
  "Поведение, которое в классическом подходе наблюдают при динамическом анализе (RegShot, Process Monitor), "
  "здесь восстановлено по машинному коду вплоть до конкретных аргументов вызовов CopyFileA, DeleteFileA и "
  "RegSetValueExA.")

# ================================================================ глава 1

H1("1 ТЕОРЕТИЧЕСКИЕ ОСНОВЫ АНАЛИЗА ВЗАИМОДЕЙСТВИЯ ВПО С ФАЙЛОВОЙ СИСТЕМОЙ")
H2("1.1 Файловая система как объект воздействия вредоносных программ")
P("Прикладная программа в Windows обращается к файловой системе через функции библиотеки kernel32.dll "
  "(Windows API). Они скрывают от программы устройство конкретной файловой системы (NTFS, FAT32, exFAT) и "
  "переводят запросы в системные вызовы ядра. Для анализа важно, что набор этих функций конечен и хорошо "
  "документирован: зная, какие из них вызывает программа и с какими аргументами, можно точно описать её "
  "воздействие на файлы.")
P(f"Функции удобно сгруппировать по видам операций (таблица {next_tab()}). Каждая группа сама по себе "
  "нейтральна — её использует любой файловый менеджер, — но характерные сочетания выдают вредоносное поведение.")
TABLE("Операции с файловой системой и их использование вредоносными программами",
      ["Операция", "Функции Windows API", "Типичное вредоносное применение"],
      [["Перечисление файлов и каталогов", "FindFirstFile, FindNextFile, FindClose", "поиск документов для кражи, шифрования или уничтожения"],
       ["Перечисление дисков", "GetLogicalDrives, GetLogicalDriveStrings, GetDriveType", "распространение действия на съёмные и сетевые диски"],
       ["Чтение и запись", "CreateFile, ReadFile, WriteFile, SetEndOfFile", "шифрование содержимого, запись записки о выкупе"],
       ["Копирование", "CopyFile, CopyFileEx", "самокопирование, затирание файлов копией шаблона"],
       ["Перемещение и замена", "MoveFile, MoveFileEx, ReplaceFile", "подмена системных файлов, отложенная замена при перезагрузке"],
       ["Удаление", "DeleteFile, RemoveDirectory, SHFileOperation", "уничтожение данных, удаление следов"],
       ["Атрибуты", "SetFileAttributes, SetFileTime", "снятие «только чтение», сокрытие файлов, подделка дат"],
       ["Системные и временные каталоги", "GetSystemDirectory, GetWindowsDirectory, GetTempPath, GetTempFileName", "размещение копии под видом системного файла, промежуточные файлы"],
       ["Собственный путь", "GetModuleFileName, GetCommandLine", "определение исходного файла для самокопирования"],
       ["Запуск файлов", "WinExec, ShellExecute, CreateProcess", "запуск полезной нагрузки или отвлекающей программы"]],
      [4.0, 5.7, 6.8])
P("Операции с файловой системой почти всегда соседствуют с операциями над реестром. Ключи автозапуска "
  "(HKLM и HKCU \\Software\\Microsoft\\Windows\\CurrentVersion\\Run) превращают скопированный файл в постоянно "
  "запускаемый, а ключ Explorer\\Shell Folders сообщает программе, где лежат «Мои документы» и рабочий стол "
  "пользователя. Поэтому модуль анализа в ИАС рассматривает реестровые функции вместе с файловыми.")

H2("1.2 Классификация техник по матрице MITRE ATT&CK")
P("Для единообразного описания поведения ВПО в работе используется общепризнанная база знаний MITRE ATT&CK, "
  f"где каждая техника имеет идентификатор. Техники, связанные с файловой системой, приведены в таблице {next_tab()}. "
  "ИАС указывает идентификатор техники у каждой найденной цепочки поведения, что позволяет сопоставлять "
  "отчёт с правилами средств защиты и с отчётами других аналитиков.")
TABLE("Техники MITRE ATT&CK, связанные с файловой системой",
      ["Идентификатор", "Название", "Проявление в коде"],
      [["T1083", "File and Directory Discovery", "FindFirstFile / FindNextFile, чтение Shell Folders, перебор дисков"],
       ["T1485", "Data Destruction", "удаление или затирание найденных файлов"],
       ["T1486", "Data Encrypted for Impact", "чтение, шифрование и перезапись файлов, записка о выкупе"],
       ["T1490", "Inhibit System Recovery", "vssadmin delete shadows, wbadmin, bcdedit"],
       ["T1222.001", "File and Directory Permissions Modification", "SetFileAttributes, изменение ACL"],
       ["T1547.001", "Registry Run Keys / Startup Folder", "RegSetValueEx в ключ Run, копирование в «Автозагрузку»"],
       ["T1036.005", "Masquerading: Match Legitimate Name or Location", "копия в System32 под системным именем"],
       ["T1564.001", "Hide Artifacts: Hidden Files and Directories", "атрибуты «скрытый» и «системный»"],
       ["T1491.001", "Defacement: Internal Defacement", "замена содержимого файлов посланием злоумышленника"],
       ["T1480", "Execution Guardrails", "именованный мьютекс, проверка окружения перед запуском"]],
      [2.8, 5.6, 8.1])

H2("1.3 Методы анализа исполняемых файлов")
P("Исполняемый файл исследуют тремя группами методов.")
P("**Статический анализ** изучает файл без запуска: заголовки PE, секции, таблицу импорта, строки, "
  "дизассемблированный код. Он безопасен и даёт полную картину всех ветвей программы, включая те, что при "
  "конкретном запуске не выполняются. Ограничения — упаковка и шифрование кода, динамическое получение адресов "
  "функций, трудоёмкость ручного разбора.")
P("**Эвристический (сигнатурный) анализ** сопоставляет найденные признаки с базой знаний: именами опасных "
  "функций, ключами реестра, характерными строками. Так работают антивирусные движки и ИАС из ПР №7. Метод "
  "быстрый, но даёт ложные срабатывания: функция DeleteFile есть и в любом установщике.")
P("**Динамический анализ** запускает программу в изолированной среде и фиксирует её действия (Process Monitor, "
  "RegShot, сетевой анализатор). Для деструктивного образца он требует полностью изолированной виртуальной "
  "машины без общих папок и сетевых дисков: программа, уничтожающая файлы на всех доступных дисках, может "
  "повредить данные хост-системы через подключённые ресурсы.")
P("В работе выбран путь, совмещающий достоинства первых двух методов: ИАС выполняет эвристический анализ строк "
  "и импорта, а затем разбирает машинный код и восстанавливает фактические аргументы вызовов. Это даёт ту же "
  "информацию, что и журнал Process Monitor (какой файл, какая операция, с какими параметрами), но без запуска "
  "образца и для всех ветвей программы сразу.")

H2("1.4 Формат PE и вызов функций через таблицу импорта")
P("Исполняемые файлы Windows имеют формат PE (Portable Executable). Файл начинается с DOS-заголовка "
  "с сигнатурой «MZ», поле e_lfanew по смещению 0x3C указывает на заголовок PE. За сигнатурой «PE\\0\\0» "
  "следуют файловый заголовок (архитектура, число секций, дата сборки) и дополнительный заголовок (точка входа, "
  "база образа, каталог данных). Каталог данных содержит адреса таблиц импорта, ресурсов, перемещений и др.")
P("Таблица импорта перечисляет библиотеки и функции, которые программа использует. При загрузке Windows "
  "записывает адреса функций в таблицу адресов импорта (IAT), а код вызывает их косвенно — командой "
  "call dword ptr [адрес ячейки IAT] (байты FF 15 и 4-байтовый адрес) либо загружает адрес в регистр "
  "(mov ebx, [адрес ячейки]) и затем выполняет call ebx. Это ключевое свойство для анализа: зная адреса "
  "ячеек IAT, можно найти в машинном коде все места вызова каждой функции. Аргументы функций в 32-битном коде "
  "передаются через стек командами push в обратном порядке, в 64-битном — через регистры rcx, rdx, r8, r9 и стек.")

# ================================================================ глава 2

H1("2 РАЗРАБОТКА ИНФОРМАЦИОННО-АНАЛИТИЧЕСКОЙ СИСТЕМЫ")
H2("2.1 Постановка задачи")
P("Техническое задание ПР №8 требует создать прототип ИАС, которая по пути к исполняемому файлу формирует "
  "отчёт с блоками «Объект анализа», «Признаки уязвимостей», «Индикаторы компрометации», «Оценка риска» и "
  "«Рекомендации». Для курсовой работы задание расширено пятым блоком «Взаимодействие с файловой системой». "
  f"Соответствие требований и реализации показано в таблице {next_tab()}.")
TABLE("Требования к ИАС и их реализация",
      ["Требование", "Реализация"],
      [["Приём пути к файлу (.exe, .dll, .bin); ошибка, если файла нет", "модуль input; дополнительно режим «-» — чтение из стандартного ввода"],
       ["Извлечение строк утилитой strings или внутренними методами", "strings.exe (Sysinternals), если найдена, иначе встроенное извлечение ASCII и UTF-16LE"],
       ["Нормализация данных", "чистка непечатаемых символов, отсев мусора, удаление повторов"],
       ["Классификация: мета-данные, ошибки, IOC, CWE/CVE, вероятность инцидента", "модуль signatures: матрица из 55 правил, база CVE, проверка структуры PE"],
       ["Анализ взаимодействия с файловой системой (тема курсовой)", "модули disasm и fsanalysis: вызовы WinAPI с аргументами, функции, цепочки поведения"],
       ["Структурированный отчёт с выводами и рекомендациями", "модуль report: шесть блоков, таблица CSV, список IOC для блокировки"],
       ["Высокоуровневый язык, работа с файлами и внешними утилитами", "C++20, Microsoft Visual Studio 18 (MSVC), Windows API (BCrypt, _popen)"]],
      [7.5, 9.0])

H2("2.2 Архитектура системы")
P(f"ИАС построена из восьми модулей (рисунок {next_fig()}, таблица {next_tab()}). Каждый модуль оформлен "
  "парой файлов .h/.cpp, общие типы данных собраны в common.h. Модули передают друг другу данные через "
  "структуру Analysis, которая в конце целиком поступает генератору отчёта.")
FIGURE("fig01_arch.png", "Архитектура ИАС: модули и потоки данных")
loc = {}
for fn in os.listdir(os.path.join(PROJ, "ExeIAS")):
    if fn.endswith((".cpp", ".h")):
        loc[fn] = sum(1 for _ in open(os.path.join(PROJ, "ExeIAS", fn), encoding="utf-8"))
TABLE("Модули ИАС",
      ["Модуль (файлы)", "Строк кода", "Назначение"],
      [["input (.h/.cpp)", loc["input.h"] + loc["input.cpp"], "приём файла, хэши MD5 / SHA-1 / SHA-256, извлечение и нормализация строк (раздел 2.1 ТЗ)"],
       ["pe", loc["pe.h"] + loc["pe.cpp"], "разбор заголовков PE, секций, энтропии, таблицы импорта"],
       ["disasm", loc["disasm.h"] + loc["disasm.cpp"], "декодер инструкций x86 / x64: длина, ModRM, SIB, смещения, операнды"],
       ["fsanalysis", loc["fsanalysis.h"] + loc["fsanalysis.cpp"], "анализ взаимодействия с ФС: вызовы и их аргументы, функции, граф вызовов, цепочки поведения"],
       ["signatures", loc["signatures.h"] + loc["signatures.cpp"], "матрица сигнатур, CWE / CVE / ATT&CK, мета-данные, ошибки, аномалии PE (раздел 2.2 ТЗ)"],
       ["report", loc["report.h"] + loc["report.cpp"], "оценка риска, отчёт TXT и таблица CSV (раздел 2.3 ТЗ)"],
       ["common", loc["common.h"] + loc["common.cpp"], "общие структуры данных и вспомогательные функции"],
       ["main", loc["main.cpp"], "разбор аргументов командной строки, порядок работы модулей"]],
      [3.6, 2.2, 10.7])
P(f"Всего {sum(loc.values())} строк на C++20. Проект собирается в Microsoft Visual Studio (решение ExeIAS.slnx, "
  "платформы x86 и x64) или из командной строки сценарием build.bat. Обе сборки проверены на одних и тех же "
  "файлах и дают побайтно одинаковые отчёты.")

H2("2.3 Приём и нормализация данных")
P("Программа запускается командой ExeIAS.exe <путь к файлу>. Если файл не существует, работа завершается "
  "с сообщением об ошибке и кодом возврата 1, как требует ТЗ. Для образцов ВПО добавлен безопасный режим: "
  "вместо пути передаётся «-», и содержимое читается из стандартного ввода. Это позволяет анализировать файл "
  "прямо из архива без записи на диск:")
LISTING("Анализ образца без распаковки архива",
        ["tar -xOf Whiter.a.rar Whiter.a.exe | ExeIAS.exe - --name Whiter.a.exe"])
P("Для файла вычисляются хэши MD5, SHA-1 и SHA-256 средствами Windows CNG (BCrypt). При чтении временно "
  "отключается перенаправление файловой системы WOW64: без этого 32-битная сборка ИАС, получив путь "
  "C:\\Windows\\System32\\cmd.exe, молча прочитала бы другой файл — из SysWOW64. Ошибка была обнаружена при "
  "сравнении отчётов 32- и 64-битной сборок и исправлена.")
P("Строки извлекаются утилитой strings.exe (Sysinternals), если она лежит рядом с программой или в PATH, "
  "иначе встроенным методом — поиском последовательностей печатных символов ASCII и UTF-16LE длиной от "
  "четырёх символов. Затем выполняется нормализация из ПР №7: удаление непечатаемых символов и пробелов по "
  "краям, отсев строк без трёх букв подряд, удаление повторов.")

H2("2.4 Матрица сигнатур")
P("Классификация строк и импорта выполняется по расширенной матрице из 55 правил. В отличие от ПР №7, где "
  "словарь содержал в основном ошибки и функции загрузки библиотек, матрица охватывает ключи реестра "
  "(Run, Winlogon, Image File Execution Options, Shell Folders), маски файлов, команды уничтожения резервных "
  "копий (vssadmin, wbadmin, bcdedit, cipher /w), фразы записок злоумышленника, сетевые функции. Каждое "
  f"правило связывает сигнатуру с категорией, CWE, техникой ATT&CK и рекомендацией (листинг {next_lst()}).")
LISTING("Структура правила матрицы сигнатур (signatures.cpp)",
        src(os.path.join(PROJ, "ExeIAS", "signatures.cpp"), "struct SignatureRule", "};") +
        ["", "    { \"CurrentVersion\\\\Run\", false, CAT_REGISTRY, \"CWE-506\", \"Embedded Malicious Code\", \"T1547.001\",",
         "      \"Закрепление в системе через ключ автозапуска\", HIGH, \"\",",
         "      \"Проверить ключи Run/RunOnce (HKLM и HKCU), удалить посторонние параметры\" },"])
P("Категория определяет блок отчёта: ошибки, небезопасные функции, уязвимые библиотеки и отсутствие защитных "
  "механизмов PE (ASLR, DEP) попадают в «Признаки уязвимостей», остальное — в «Индикаторы компрометации». "
  "Отдельно извлекаются мета-данные: имена DLL и файлов, пути, номера версий, сведения о разработчике и "
  "компиляторе. Сохранён механизм контекста из ПР №7: признаки загрузчика .NET (apphost, hostfxr) понижают "
  "критичность находок, которые для такого загрузчика штатны.")
P("Правила для фраз-посланий (piracy, decrypt, your files) проверяются только по строкам с пробелами и не "
  "проверяются по таблице импорта. Без этого ограничения слово decrypt срабатывало на штатную функцию "
  "шифрования EFS DecryptFileW в notepad.exe — такие случаи выявлялись проверкой на легитимных программах.")

H2("2.5 Модуль анализа взаимодействия с файловой системой")
P(f"Модуль fsanalysis — основная часть работы. Его конвейер показан на рисунке {next_fig()}.")
FIGURE("fig02_pipeline.png", "Этапы анализа машинного кода в модуле fsanalysis")
H3("2.5.1 Поиск обращений к функциям и строкам")
P("Сначала все исполняемые секции просматриваются побайтно в поисках шаблонов, ссылающихся на ячейки IAT и "
  "на строки в секциях данных: call [IAT] и jmp [IAT] (FF 15, FF 25), mov регистр, [адрес] (8B /r, A1), "
  "push адрес и mov регистр, адрес (68, B8–BF), для x64 — lea и mov с адресацией относительно rip. "
  "Адрес принимается, только если он точно совпадает с ячейкой IAT или попадает в известную строку, поэтому "
  "случайные совпадения байтов практически исключены. Здесь же отмечаются команды сравнения байтового регистра "
  "с символом «z» — признак цикла перебора букв дисков.")
H3("2.5.2 Границы функций и граф вызовов")
P("Начала функций определяются по целям команд call rel32 (E8): цель принимается, если перед ней стоит "
  "выравнивание (nop, int3) или конец предыдущей функции (ret), либо если она похожа на типичный пролог MSVC "
  "(push ebp / mov ebp, esp; sub esp, N) и вызывается не менее двух раз. Каждое обращение к API и строке "
  "относится к функции, в которую попадает его адрес. Вызов функцией самой себя отмечается как рекурсия, "
  "вызовы между функциями образуют граф.")
H3("2.5.3 Декодирование инструкций и восстановление аргументов")
P("Чтобы узнать не только факт вызова, но и его аргументы, каждая функция декодируется полностью. Модуль disasm "
  "реализует декодер длины инструкций x86 и x64: префиксы, REX, однобайтовые и двухбайтовые коды операций, "
  "байты ModRM и SIB, смещения и непосредственные операнды. Мнемоники не формируются — модулю анализа "
  "достаточно кода операции и операндов.")
P("Затем инструкции функции «проигрываются» на модели процессора, в которой вместо чисел хранятся смысловые "
  f"значения (таблица {next_tab()}). Модель отслеживает содержимое шестнадцати регистров, стек положенных "
  "аргументов, глубину стека относительно входа в функцию, ячейки кадра стека и содержимое буферов.")
TABLE("Виды значений в модели процессора",
      ["Вид", "Пример в отчёте", "Откуда берётся"],
      [["Константа", "FALSE, 0x100, HKEY_LOCAL_MACHINE", "push imm, mov reg, imm, xor reg, reg"],
       ["Строка образа", "\"c:\\wxp\"", "адрес, попадающий в строку секции данных"],
       ["Функция API", "CopyFileA", "mov reg, [ячейка IAT] перед call reg"],
       ["Буфер в стеке", "var_444 = \"<arg_0>\\<найденный файл>\"", "lea reg, [esp+N]; содержимое — из копирования строк и результатов API"],
       ["Аргумент функции", "arg_0", "чтение [esp+N] выше адреса возврата"],
       ["Символ", "hKey(HKLM\\SOFTWARE\\...\\Run), FILE*(c:\\wxp)", "результат API или функции fopen"]],
      [3.4, 6.4, 6.7])
P("Особенности, которые пришлось учесть для компилятора MSVC 6.0, которым собран образец:")
L(["соглашение stdcall: функция WinAPI сама снимает аргументы со стека, их число берётся из таблицы "
   "прототипов (91 функция); для внутренних функций — из команды ret N в их коде",
   "соглашение cdecl: стек чистит вызывающий командой add esp, N, причём компилятор объединяет очистку после "
   "нескольких вызовов; число аргументов вычисляется с учётом уже отнесённых к предыдущему вызову значений",
   "встроенные strcpy и strcat: последовательность repne scasb / not ecx / sub edi, ecx / rep movsd "
   "распознаётся как копирование или дописывание строки в буфер",
   "результаты API: GetSystemDirectoryA записывает в буфер «%SystemRoot%\\system32», GetTempFileNameA — имя "
   "временного файла, RegCreateKeyA — описатель ключа с полным путём, FindFirstFileA — имя найденного файла "
   "в поле cFileName структуры WIN32_FIND_DATA",
   "переходы: в точке условного перехода запоминается состояние регистров и стека, после ret или jmp анализ "
   "продолжается с сохранённым состоянием",
   "sprintf-подобные функции (буфер, формат, …) — подстановка аргументов в формат"])
P(f"Фрагмент обработки вызова функции WinAPI приведён в листинге {next_lst()}.")
LISTING("Восстановление аргументов вызова WinAPI (fsanalysis.cpp)",
        src(os.path.join(PROJ, "ExeIAS", "fsanalysis.cpp"), "    void apiCall(const string& name, size_t idx)",
            "        f.calls.push_back(cs);"), style="CodeSmall")
H3("2.5.4 Цепочки поведения")
P("По восстановленным вызовам каждая функция проверяется на сочетания операций, характерные для вредоносного "
  f"поведения (таблица {next_tab()}). Сочетание внутри одной функции и, тем более, связка функций по графу "
  "вызовов значительно надёжнее отдельного имени функции в таблице импорта.")
TABLE("Правила выявления цепочек поведения",
      ["Цепочка", "Условие в коде функции", "ATT&CK", "Критичность"],
      [["Обход каталогов", "FindFirstFile + FindNextFile (рекурсивный, если функция вызывает себя)", "T1083", "LOW"],
       ["Массовое уничтожение файлов", "обход + DeleteFile / RemoveDirectory", "T1485", "CRITICAL"],
       ["Массовая перезапись", "обход + CopyFile / WriteFile / MoveFile без удаления", "T1565.001", "HIGH"],
       ["Снятие защитных атрибутов", "обход + SetFileAttributes (HIGH при уничтожении или перезаписи)", "T1222.001", "MEDIUM / HIGH"],
       ["Самокопирование в системный каталог", "GetSystemDirectory + CopyFile + GetCommandLine / GetModuleFileName", "T1036.005", "HIGH"],
       ["Закрепление через автозапуск", "ключ …\\CurrentVersion\\Run + RegSetValueEx", "T1547.001", "HIGH"],
       ["Поиск каталогов пользователя", "ключ Shell Folders + RegQueryValueEx", "T1083", "MEDIUM"],
       ["Перебор логических дисков", "GetLogicalDrives / GetDriveType или цикл по буквам до «z»", "T1083", "MEDIUM"],
       ["Файл с текстом-посланием", "fopen-подобный вызов + запись строки в этот файл", "T1491.001", "MEDIUM"],
       ["Временный файл и запуск программы", "GetTempPath / GetTempFileName + WinExec / CreateProcess", "T1106", "MEDIUM"],
       ["Контроль единственного экземпляра", "CreateMutex + GetLastError", "T1480", "LOW"],
       ["Целевое уничтожение документов", "функция с поиском каталогов пользователя вызывает функцию уничтожения", "T1485", "CRITICAL"],
       ["Уничтожение на всех дисках", "функция с перебором дисков вызывает функцию уничтожения", "T1485", "CRITICAL"]],
      [4.2, 7.0, 2.3, 3.0], size=11)
P("В доказательства каждой цепочки включаются восстановленные вызовы с адресами и пояснениями смысла "
  "аргументов, например «файл-приёмник перезаписывается, если существует» для CopyFile с bFailIfExists = FALSE "
  "или «автозапуск для всех пользователей (нужны права администратора)» для ключа Run в разделе HKLM.")

H2("2.6 Оценка риска и генерация отчёта")
P("Каждая находка получает уровень критичности LOW, MEDIUM, HIGH или CRITICAL с весами 1, 3, 8 и 20. "
  "Цепочки, подтверждённые машинным кодом, весят вдвое больше находок по строкам и импорту: функции двойного "
  "назначения есть и в легитимных программах, а сочетание операций в коде — гораздо более сильный признак. "
  "Однотипная цепочка, найденная в нескольких функциях, учитывается один раз. Сумма ограничивается 100 баллами "
  "и переводится в вероятность инцидента ИБ: менее 20 — низкая, 20–44 — средняя, 45–69 — высокая, 70 и более — "
  "очень высокая. При признаках легитимного загрузчика .NET и отсутствии цепочек балл уменьшается вдвое.")
P("Итоговый вывод формируется правилами: наличие кода уничтожения данных — «вредоносная программа "
  "деструктивного типа»; подтверждённая кодом цепочка уровня HIGH или высокий балл — «требуется углублённый "
  "анализ»; только находки по строкам — «функции двойного назначения, вредоносность не установлена».")
P("Отчёт report_<имя>.txt состоит из шести блоков: «Объект анализа» (хэши, заголовки, секции, импорт, "
  "мета-данные), «Признаки уязвимостей» (ошибки, CWE, CVE), «Индикаторы компрометации», «Взаимодействие "
  "с файловой системой» (операции, карта функций с восстановленными вызовами, граф вызовов, цепочки), "
  "«Оценка риска» и «Рекомендации» со списком индикаторов для блокировки. Все находки дублируются в таблицу "
  "report_<имя>.csv для обработки в Excel.")

# ================================================================ глава 3

H1("3 АНАЛИЗ ОБРАЗЦА WHITER.A")
H2("3.1 Объект анализа и меры безопасности")
P("Образец получен в виде архива Whiter.a.rar (21 282 байта), содержащего единственный файл Whiter.a.exe "
  f"размером 57 344 байта. Хэши файла приведены в таблице {next_tab()}.")
TABLE("Идентификаторы образца", ["Алгоритм", "Значение"],
      [["MD5", "799b57227561238a7d7a284c5568c1ad"],
       ["SHA-1", "f62ddd138ab15b67a2207438b38414fd236d5278"],
       ["SHA-256", "fe974c995cfb27e8c91123081986847f6d3d4252b6a8d1e1385c558f2aeb7057"]],
      [3.0, 13.5])
P("Файл не запускался и не записывался на диск в распакованном виде: при распаковке его, вероятнее всего, "
  "удалил бы антивирус, а случайный запуск недопустим. Все исследования выполнены по содержимому, переданному "
  "из архива через стандартный ввод (листинг 1). Для проверки результатов ИАС использовалось независимое "
  "дизассемблирование тех же байтов библиотекой Capstone — адреса и вызовы совпали полностью.")

H2("3.2 Структура PE-файла")
P(f"DOS-заголовок образца стандартный (рисунок {next_fig()}): сигнатура «MZ», поле e_lfanew = 0xE0, "
  "заглушка с текстом «This program cannot be run in DOS mode.». Между заглушкой и заголовком PE расположен "
  "заголовок Rich — служебная запись компоновщика Microsoft о версиях использованных инструментов.")
FIGURE("fig03_dos.png", "Шестнадцатеричное представление DOS-заголовка и заглушки")
P(f"Заголовок PE и основные поля показаны на рисунке {next_fig()} и в таблице {next_tab()}.")
FIGURE("fig04_pe.png", "Шестнадцатеричное представление заголовка PE")
TABLE("Основные поля заголовка PE", ["Поле", "Значение", "Вывод"],
      [["Machine", "0x014C", "32-битная программа x86"],
       ["NumberOfSections", "4", ".text, .rdata, .data, .rsrc"],
       ["TimeDateStamp", "0x3B874509", "25.08.2001 06:26:17 UTC — эпоха Windows XP (кодовое имя Whistler)"],
       ["Characteristics", "0x010F", "исполняемый файл, перемещения и отладочные сведения удалены"],
       ["Magic / Linker", "0x10B / 6.0", "PE32, компоновщик Microsoft Visual C++ 6.0"],
       ["AddressOfEntryPoint", "0x1995", "точка входа в секции .text (код запуска CRT)"],
       ["ImageBase", "0x00400000", "стандартная база для EXE"],
       ["Subsystem", "2", "графическое приложение — окно консоли не появляется"],
       ["DllCharacteristics", "0x0000", "нет ASLR и DEP (норма для 2001 года)"],
       ["DataDirectory[1] Import", "RVA 0x9580, 0x3C", "две записи таблицы импорта"],
       ["DataDirectory[2] Resource", "RVA 0xF000, 0x3A0", "ресурсы: только значок"],
       ["CheckSum", "0", "контрольная сумма не вычислялась"]],
      [4.4, 3.6, 8.5], size=12)
P(f"Таблица секций (рисунок {next_fig()}, таблица {next_tab()}) не содержит аномалий: энтропия кода 6,19 "
  "типична для неупакованного машинного кода (у упакованного она выше 7,2), виртуальные и файловые размеры "
  "согласованы, точка входа лежит в исполняемой секции, оверлея нет. Следовательно, образец не упакован и "
  "доступен для статического анализа без предварительной распаковки.")
FIGURE("fig05_sections.png", "Шестнадцатеричное представление заголовков секций")
TABLE("Секции образца", ["Секция", "Виртуальный адрес", "VirtualSize", "SizeOfRawData", "Права", "Энтропия"],
      [[".text", "0x1000", "0x7044", "0x8000", "R-X", "6,19"],
       [".rdata", "0x9000", "0x0B82", "0x1000", "R--", "4,36"],
       [".data", "0xA000", "0x419C", "0x3000", "RW-", "0,81"],
       [".rsrc", "0xF000", "0x03A0", "0x1000", "R--", "0,85"]],
      [2.2, 3.0, 2.8, 3.0, 2.0, 2.4])
P("Секция ресурсов содержит только значок программы (RT_ICON, 744 байта, и группа значков RT_GROUP_ICON) "
  "с кодом языка 1041 (японский). Ни окон, ни строковых таблиц нет — интерфейс программы ограничивается "
  "Блокнотом, который она запускает сама.")

H2("3.3 Таблица импорта")
P(f"Образец импортирует 68 функций из двух библиотек: 64 из KERNEL32.dll и 4 из ADVAPI32.dll. Функции, "
  f"относящиеся к теме работы, сгруппированы в таблице {next_tab()}.")
TABLE("Импорт образца по группам", ["Группа", "Функции"],
      [["Перечисление файлов", "FindFirstFileA, FindNextFileA, FindClose"],
       ["Копирование, удаление, атрибуты", "CopyFileA, DeleteFileA, SetFileAttributesA"],
       ["Каталоги и временные файлы", "GetSystemDirectoryA, GetTempPathA, GetTempFileNameA"],
       ["Собственный путь", "GetCommandLineA, GetModuleFileNameA"],
       ["Ввод-вывод (используется CRT: fopen, fprintf)", "CreateFileA, ReadFile, WriteFile, SetFilePointer, SetEndOfFile, FlushFileBuffers"],
       ["Запуск и синхронизация", "WinExec, CreateMutexA, GetLastError"],
       ["Реестр (ADVAPI32.dll)", "RegCreateKeyA, RegSetValueExA, RegQueryValueExA, RegCloseKey"],
       ["Остальные (библиотека CRT)", "память (HeapAlloc, VirtualAlloc), время (GetLocalTime, GetSystemTime), "
        "окружение, кодовые страницы, LoadLibraryA / GetProcAddress (для MessageBoxA при ошибке CRT)"]],
      [5.2, 11.3], size=12)
P("Уже по импорту видно сочетание, нехарактерное для обычной программы: перечисление файлов вместе с "
  "копированием, удалением и снятием атрибутов, запись в реестр и запуск программ. Сетевых функций нет — "
  "образец действует автономно и не связывается с управляющим сервером. Скрытия импорта (малое число "
  "функций при наличии GetProcAddress) также нет: LoadLibraryA и GetProcAddress используются только "
  "стандартной библиотекой для вывода сообщения об ошибке.")

H2("3.4 Анализ строк")
P(f"ИАС извлекла 274 строки, после нормализации осталось 168 уникальных. Большая часть — сообщения "
  "библиотеки времени выполнения Microsoft Visual C++ («runtime error», «not enough space for "
  "environment» и др.) и имена импортируемых функций. Все содержательные строки программы сосредоточены "
  f"в секции .data (рисунок {next_fig()}, таблица {next_tab()}).")
FIGURE("fig06_data.png", "Строки программы в секции .data")
TABLE("Значимые строки образца", ["Смещение", "Строка", "Назначение (по анализу кода)"],
      [["0xA040", "c:\\wxp", "файл-шаблон, копией которого затираются файлы"],
       ["0xA04C, 0xA054", "%s\\%s, \\*.*", "сборка пути к найденному файлу, маска перебора"],
       ["0xA05C", "notepad.exe %s", "показ временного файла в Блокноте"],
       ["0xA070, 0xA074", "~sn, 0123…XYZ", "префикс временного файла и алфавит «ключа»"],
       ["0xA09C", "c:\\zwxp", "пробный файл для проверки диска на запись"],
       ["0xA0A8, 0xA0B4", "Personal, …\\Explorer\\Shell Folders", "путь к папке «Мои документы»"],
       ["0xA0F8", "You did a piracy, you deserve it.", "послание, записываемое в файл-шаблон"],
       ["0xA11C, 0xA128", "Whistler, …\\CurrentVersion\\Run", "имя параметра и ключ автозапуска"],
       ["0xA158, 0xA160", " -next, \\whismng.exe", "аргумент второго запуска, имя копии"],
       ["0xA170", "WhistlerMutex", "мьютекс единственного экземпляра"]],
      [3.0, 6.2, 7.3], size=12)
P("«Whistler» — кодовое имя Windows XP в период разработки. Вместе с текстом «You did a piracy, you deserve it» "
  "(«Ты занимался пиратством — ты это заслужил») и строкой-алфавитом это позволяет предположить маскировку "
  "под генератор ключей для нелицензионной Windows XP. Анализ кода подтверждает это предположение.")

H2("3.5 Результаты работы ИАС")
P(f"Вывод ИАС в консоли показан на рисунке {next_fig()}. Система разобрала все 96 найденных функций, "
  "восстановила аргументы 264 вызовов и выявила 12 цепочек поведения. Итог — 25 находок, уровень риска "
  "CRITICAL, балл 100 из 100, вероятность инцидента «очень высокая», вывод — «вредоносная программа "
  "деструктивного типа (троян-вайпер)».")
FIGURE("fig07_console.png", "Результат работы ИАС для образца Whiter.a")
TABLE("Находки ИАС по блокам отчёта", ["Блок", "Находки"],
      [["Признаки уязвимостей (4)", "CWE-426 (загрузка KERNEL32.dll), CWE-248 (abnormal program termination), "
        "CWE-693: нет ASLR, нет DEP"],
       ["Индикаторы компрометации (9)", "WinExec; ключ CurrentVersion\\Run; LoadLibraryA; Shell Folders; "
        "RegSetValueEx; текст-послание; файлы в корне диска c:\\wxp и c:\\zwxp; имя \\whismng.exe; notepad.exe"],
       ["Взаимодействие с ФС (12 цепочек)", "CRITICAL: массовое уничтожение файлов, целевое уничтожение "
        "документов, уничтожение на всех дисках; HIGH: снятие атрибутов, самокопирование, автозапуск; MEDIUM: "
        "временный файл и запуск программы, файл с посланием, поиск каталогов пользователя, перебор дисков; "
        "LOW: рекурсивный обход, мьютекс"],
       ["Оценка риска", "CRITICAL — 3, HIGH — 5, MEDIUM — 11, LOW — 6; балл 100/100"]],
      [5.0, 11.5], size=12)
P(f"Карта функции уничтожения из отчёта ИАС приведена в листинге {next_lst()}. В ней видны не только имена "
  "функций, но и восстановленные аргументы: путь к файлу собирается из аргумента функции и имени найденного "
  "файла, копируется шаблон c:\\wxp, затем удаляется тот же файл.")
LISTING("Фрагмент отчёта ИАС: функция sub_401000",
        report_block("  Функция sub_401000", "  Функция sub_4011E0")[:-1], style="CodeSmall")

H2("3.6 Восстановление алгоритма работы образца")
P(f"По результатам ИАС и ручного разбора дизассемблированного кода восстановлен полный алгоритм "
  f"(рисунок {next_fig()}). Пользовательский код занимает три функции по адресам 0x401000–0x40164F, "
  f"остальной код образца — стандартная библиотека Visual C++ 6.0. Граф вызовов показан на рисунке {next_fig() + 1}.")
FIGURE("fig08_algorithm.png", "Схема алгоритма работы Whiter.a", width=15.5)
FIGURE("fig10_callgraph.png", "Граф вызовов функций образца и их роли")

H3("3.6.1 Выбор режима работы")
P(f"Главная функция sub_401300 (WinMain) сначала создаёт мьютекс «WhistlerMutex» (листинг {next_lst()}). "
  "Если GetLastError возвращает 0xB7 (ERROR_ALREADY_EXISTS), значит, копия программы уже работает, и "
  "функция завершается. Затем проверяется длина командной строки (аргумент lpCmdLine): при первом запуске она "
  "пуста, при автозапуске содержит «-next».")
LISTING("Создание мьютекса и выбор режима (sub_401300)",
        asm_lines(0x401300, 0x40134E, {0x40130F: "CreateMutexA(NULL, TRUE, \"WhistlerMutex\")",
                                       0x40131B: "ERROR_ALREADY_EXISTS?",
                                       0x401331: "edi = lpCmdLine",
                                       0x40134E: "аргументы есть -> разрушение"}))
H3("3.6.2 Первый запуск: закрепление и отвлекающий манёвр")
P(f"При первом запуске программа определяет свой путь по командной строке (кавычки отбрасываются вызовом "
  f"strchr), формирует путь «%SystemRoot%\\system32\\whismng.exe» и копирует себя туда (листинг {next_lst()}). "
  "Имя whismng.exe подобрано так, чтобы выглядеть системным компонентом («Whistler manager»). Если "
  "копирование не удалось, управление сразу переходит к разрушительной части.")
LISTING("Самокопирование и запись ключа автозапуска (sub_401300)",
        asm_lines(0x4013D8, 0x401430, {0x4013E0: "GetSystemDirectoryA(var_100, 100h)",
                                       0x4013FD: "bFailIfExists = FALSE (eax = 0)",
                                       0x401428: "CopyFileA(<свой путь>, ...\\whismng.exe, FALSE)",
                                       0x401430: "не скопировано -> разрушение"}) + ["..."] +
        asm_lines(0x40144C, 0x40144C, {0x40144C: "80000002h = HKEY_LOCAL_MACHINE"}) + ["..."] +
        asm_lines(0x4014C5, 0x401513, {0x4014D6: "RegCreateKeyA(HKLM, \"...\\Run\", &hKey)",
                                       0x4014FC: "RegSetValueExA(hKey, \"Whistler\", 0, REG_SZ, ...)",
                                       0x40150C: "показ «ключа» в Блокноте",
                                       0x401513: "Блокнот не запущен -> разрушение"}))
P("Далее в раздел HKEY_LOCAL_MACHINE ключа автозапуска записывается параметр «Whistler» со значением "
  "«%SystemRoot%\\system32\\whismng.exe -next». ИАС восстановила этот вызов полностью:")
LISTING("Восстановленный ИАС вызов записи в реестр",
        ["RegSetValueExA(hKey(HKLM\\SOFTWARE\\Microsoft\\Windows\\CurrentVersion\\Run), \"Whistler\", 0,",
         "               REG_SZ, var_528 = \"%SystemRoot%\\system32\\whismng.exe -next\", ?)"])
P(f"После этого вызывается функция sub_4011E0 (листинг {next_lst()}). Она инициализирует генератор случайных "
  "чисел текущим временем и составляет строку из 29 символов: позиции 5, 11, 17 и 23 заполняются дефисами, "
  "первый символ — случайной латинской буквой, остальные — случайными символами из алфавита "
  "«0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ». Получается строка формата XXXXX-XXXXX-XXXXX-XXXXX-XXXXX — "
  "точно как ключ продукта Windows XP. Строка записывается во временный файл %TEMP%\\~snXXXX.tmp, который "
  "открывается Блокнотом (WinExec с параметром SW_SHOWDEFAULT).")
LISTING("Генерация поддельного ключа и его показ (sub_4011E0)",
        asm_lines(0x40120C, 0x40125D, {0x40120E: "позиции 5, 11, 17, 23 — дефис",
                                       0x401226: "rand() % 26 — первая буква",
                                       0x40123D: "rand() % 36 — буква или цифра",
                                       0x401254: "'-'",
                                       0x40125A: "29 символов"}) + ["..."] +
        asm_lines(0x401271, 0x4012E7, {0x401271: "GetTempPathA(104h, var_208)",
                                       0x40128B: "GetTempFileNameA(%TEMP%, \"~sn\", 0, var_30C)",
                                       0x40129B: "fopen(var_30C, \"w\")",
                                       0x4012AB: "fprintf(файл, ключ)",
                                       0x4012CB: "sprintf(cmd, \"notepad.exe %s\", var_30C)",
                                       0x4012DD: "WinExec(cmd, SW_SHOWDEFAULT)",
                                       0x4012E7: "успех, если WinExec вернула > 31"}))
P("Пользователь, запустивший «генератор ключей», видит ожидаемый результат — ключ в Блокноте — и не "
  "подозревает о закреплении. Программа завершается, а разрушительная часть отложена до следующего входа в "
  "систему, когда Windows запустит whismng.exe -next из ключа автозапуска. Задержка между запуском и ущербом "
  "мешает пользователю связать одно с другим.")

H3("3.6.3 Разрушительная часть")
P(f"При запуске с аргументом (или при неудаче первых шагов) выполняется код по адресу 0x401524 "
  f"(листинг {next_lst()}). Сначала в корне диска C: создаётся файл c:\\wxp с текстом «You did a piracy, you "
  "deserve it.» — это шаблон, которым будут затираться файлы. Затем из ключа HKCU\\...\\Explorer\\Shell "
  "Folders читается параметр Personal — путь к папке «Мои документы», и он передаётся функции уничтожения.")
LISTING("Создание файла-шаблона и уничтожение «Моих документов» (sub_401300)",
        asm_lines(0x401524, 0x4015D6, {0x40152E: "fopen(\"c:\\\\wxp\", \"w\")",
                                       0x40153E: "fprintf(файл, послание)",
                                       0x401544: "fclose",
                                       0x401566: "80000001h = HKEY_CURRENT_USER",
                                       0x401589: "RegCreateKeyA(HKCU, \"...\\Shell Folders\", &hKey)",
                                       0x4015C1: "RegQueryValueExA(hKey, \"Personal\", ...)",
                                       0x4015D6: "sub_401000(<Мои документы>)"}), style="CodeSmall")
P(f"Далее (листинг {next_lst()}) программа перебирает буквы дисков от «d» (0x64) до «z» (0x7A). Для каждой "
  "буквы она подставляет её в шаблоны «c:» и «c:\\zwxp» и пытается создать пробный файл <диск>:\\zwxp. "
  "Если файл создан — диск существует и доступен для записи, и его корень передаётся функции уничтожения. "
  "Так проверяются жёсткие диски, флеш-накопители и подключённые сетевые диски; привод компакт-дисков "
  "отсеивается, потому что на него нельзя записать. Последним уничтожается диск C:, после чего команда "
  "jmp возвращает выполнение к уничтожению «Моих документов» — цикл бесконечен.")
LISTING("Перебор дисков и бесконечный цикл (sub_401300)",
        asm_lines(0x401607, 0x40164E, {0x401607: "bl = 'd'",
                                       0x401613: "буква в шаблон \"<диск>:\"",
                                       0x401617: "буква в \"<диск>:\\\\zwxp\"",
                                       0x40161B: "fopen(\"<диск>:\\\\zwxp\", \"w\")",
                                       0x401625: "не создан -> следующий диск",
                                       0x401632: "sub_401000(\"<диск>:\")",
                                       0x40163C: "до 'z'",
                                       0x401646: "sub_401000(\"c:\")",
                                       0x40164E: "снова «Мои документы» — бесконечно"}))

H3("3.6.4 Функция уничтожения sub_401000")
P(f"Функция sub_401000 принимает путь к каталогу и обходит его рекурсивно (листинг {next_lst()}, "
  f"рисунок {next_fig()}). К пути дописывается маска «\\*.*», и FindFirstFileA / FindNextFileA перебирают "
  "все элементы. Элементы, имя которых начинается с точки, пропускаются (это «.», «..», а заодно и файлы "
  "вида .gitignore). Для каждого элемента функцией sprintf собирается полный путь «<каталог>\\<имя>».")
LISTING("Функция уничтожения файлов (sub_401000)",
        asm_lines(0x401040, 0x401093, {0x401040: "\"\\\\*.*\" — дописать маску",
                                       0x401078: "FindFirstFileA(var_100, var_340)",
                                       0x40108D: "ebx = CopyFileA",
                                       0x401093: "cFileName[0] == '.' ?"}) + ["..."] +
        asm_lines(0x4010A1, 0x4010C5, {0x4010B4: "sprintf(var_444, \"%s\\\\%s\", каталог, имя)",
                                       0x4010C3: "FILE_ATTRIBUTE_DIRECTORY?"}) + ["..."] +
        asm_lines(0x401166, 0x401166, {0x401166: "рекурсия для подкаталога"}) +
        asm_lines(0x401174, 0x4011C4, {0x401174: "READONLY | SYSTEM ?",
                                       0x401182: "SetFileAttributesA(путь, FILE_ATTRIBUTE_NORMAL)",
                                       0x401194: "CopyFileA(\"c:\\\\wxp\", путь, FALSE)",
                                       0x40119B: "DeleteFileA(путь)",
                                       0x4011B5: "FindNextFileA",
                                       0x4011C4: "FindClose"}), style="CodeSmall")
P("Если элемент — каталог, функция вызывает сама себя. Для файла выполняются три операции подряд:")
L(["если установлен атрибут «только чтение» или «системный» (проверка test al, 5), атрибуты сбрасываются "
   "в FILE_ATTRIBUTE_NORMAL — иначе следующие операции завершились бы отказом",
   "CopyFileA(\"c:\\wxp\", файл, FALSE) — копия шаблона записывается поверх файла: параметр bFailIfExists = "
   "FALSE разрешает перезапись, и содержимое любого документа заменяется 33 байтами послания",
   "DeleteFileA(файл) — тот же файл удаляется, минуя корзину"], numbered=True)
FIGURE("fig09_wipe.png", "Алгоритм функции уничтожения sub_401000", width=14)
P("Результаты ни одного из вызовов не проверяются: ошибки копирования или удаления (например, у открытого "
  "другой программой файла) молча игнорируются, и обход продолжается.")

H2("3.7 Последствия для файловой системы")
P("По восстановленному алгоритму можно точно описать состояние системы после срабатывания образца:")
L(["в «Моих документах», на всех доступных для записи дисках и на диске C: удаляются все файлы, кроме "
   "занятых другими процессами и имён, начинающихся с точки",
   "структура каталогов сохраняется, но каталоги пусты: функции RemoveDirectory в образце нет",
   "атрибуты «только чтение» и «системный» не защищают файлы — они снимаются перед удалением",
   "затираются в том числе файлы операционной системы на диске C:, не занятые в момент обхода, поэтому "
   "после перезагрузки система, вероятнее всего, не загрузится",
   "уничтожаются данные на флеш-накопителях и сетевых дисках, подключённых с правом записи",
   "цикл бесконечен: файлы, созданные пользователем после заражения, тоже уничтожаются, пока процесс "
   "whismng.exe не будет остановлен; постоянный обход дисков создаёт высокую нагрузку на диск"])
P("Затирание шаблоном перед удалением не позволяет восстановить содержимое штатными средствами (из "
  "«удалённого» файла получился бы текст послания). Однако CopyFile не перезаписывает старые кластеры "
  "нулями: файл-приёмник усекается, его прежние кластеры помечаются свободными. Поэтому часть данных может "
  "быть восстановлена средствами компьютерной криминалистики, если после заражения диск больше не "
  "использовался, — важное основание для немедленного выключения узла.")
P("Интересная особенность логики: при обходе диска C: функция доходит и до самого шаблона c:\\wxp. "
  "Копирование файла в самого себя не выполнится, а удаление — выполнится. Начиная с этого момента CopyFileA "
  "завершается ошибкой (исходного файла больше нет), и файлы просто удаляются без затирания. Поскольку NTFS "
  "выдаёт элементы каталога в алфавитном порядке, большинство каталогов корня C: (Program Files, Users, "
  "Windows) обрабатывается раньше, чем файл wxp.")
P("Образец создан для Windows XP, где пользователь обычно работал с правами администратора. На современной "
  "Windows поведение отличается. Без повышения прав запись в System32 и в раздел HKLM запрещена, CopyFileA "
  "возвращает ошибку — и по логике кода управление сразу переходит к разрушительной части, без отложенного "
  "запуска. Для 32-битной программы на 64-битной Windows обращения к System32 и HKLM\\SOFTWARE перенаправляются "
  "механизмом WOW64 в SysWOW64 и WOW6432Node. Создание файла в корне диска C: обычному пользователю также "
  "запрещено, а результат fopen не проверяется (CWE-252); дальнейшее поведение в этом случае определяется "
  "реакцией библиотеки CRT на нулевой указатель. Эти выводы сделаны по коду и требуют подтверждения "
  "динамическим анализом в изолированной среде.")

H2("3.8 Индикаторы компрометации и рекомендации")
P(f"ИАС сформировала список индикаторов для поиска заражённых узлов и блокировки (таблица {next_tab()}).")
TABLE("Индикаторы компрометации Whiter.a", ["Тип", "Значение"],
      [["SHA-256 файла", "fe974c995cfb27e8c91123081986847f6d3d4252b6a8d1e1385c558f2aeb7057"],
       ["Файл", "%SystemRoot%\\system32\\whismng.exe (на 64-битной Windows — SysWOW64)"],
       ["Параметр реестра", "HKLM\\SOFTWARE\\Microsoft\\Windows\\CurrentVersion\\Run, «Whistler» = «…\\whismng.exe -next»"],
       ["Мьютекс", "WhistlerMutex"],
       ["Файлы в корне дисков", "c:\\wxp с текстом «You did a piracy, you deserve it.»; <диск>:\\zwxp"],
       ["Временные файлы", "%TEMP%\\~snXXXX.tmp со строкой вида XXXXX-XXXXX-XXXXX-XXXXX-XXXXX"],
       ["Признак в файлах", "файлы размером 33 байта с текстом послания"]],
      [4.0, 12.5], size=12)
P("Рекомендации по реагированию, если образец был запущен:", keep=True)
L(["немедленно отключить узел от сети и извлечь съёмные носители; если разрушительная часть уже работает — "
   "выключить компьютер, не дожидаясь корректного завершения работы",
   "не загружать заражённую систему; удаление выполнять с загрузочного носителя: удалить whismng.exe, "
   "параметр «Whistler» ключа Run, файлы c:\\wxp и <диск>:\\zwxp",
   "если ущерб уже нанесён, снять образ диска и попытаться восстановить данные средствами криминалистики до "
   "любой записи на диск, затем восстановить систему и данные из резервной копии",
   "проверить сетевые папки, к которым узел имел доступ на запись"])
P("Рекомендации по предотвращению:", keep=True)
L(["не использовать «генераторы ключей» и нелицензионное ПО — именно этот путь заражения эксплуатирует образец",
   "работать под учётной записью без прав администратора: тогда закрепление в System32 и HKLM невозможно",
   "включить контролируемый доступ к папкам (Controlled Folder Access) Microsoft Defender — он запрещает "
   "неизвестным программам изменять и удалять файлы в каталогах пользователя",
   "хранить резервные копии по правилу 3-2-1, в том числе офлайн-копию, недоступную с рабочего узла",
   "отслеживать в EDR и Sysmon события: запись в ключ Run (событие 13), массовое удаление файлов одним "
   "процессом, создание мьютекса WhistlerMutex"])

H2("3.9 Проверка ИАС на других файлах")
P("Для оценки ложных срабатываний ИАС применена к легитимным программам Windows, приложению .NET из ПР №3 "
  f"и демонстрационным образцам ПР №7. Результаты приведены на рисунке {next_fig()} и в таблице {next_tab()}.")
FIGURE("fig11_scores.png", "Балл риска для исследованных файлов")

def summary(fn):
    t = open(os.path.join(PROJ, "build", fn), encoding="utf-8-sig").read()
    score = re.search(r"Балл риска:\s+(\d+)", t).group(1)
    level = re.search(r"Общий уровень риска:\s+(\S+)", t).group(1)
    chains = re.search(r"4\.4 Цепочки поведения \((\d+)\)", t)
    verdict = re.search(r"^Вывод: (.+)$", t, re.M).group(1)
    return score, level, chains.group(1) if chains else "–", verdict

rows = []
for fn, title in (("report_Whiter.a.exe.txt", "Whiter.a.exe (образец ВПО)"),
                  ("report_sample_dropper.exe.txt", "sample_dropper.exe (демо ПР №7)"),
                  ("report_cmd.exe.txt", "cmd.exe (Windows, x64)"),
                  ("report_notepad.exe.txt", "notepad.exe (Windows, x64)"),
                  ("report_ConsoleApp28.exe.txt", "ConsoleApp28.exe (.NET, ПР №3)"),
                  ("report_sample_dotnet.exe.txt", "sample_dotnet.exe (демо ПР №7)"),
                  ("report_sample_clean.exe.txt", "sample_clean.exe (демо ПР №7)"),
                  ("report_xcopy.exe.txt", "xcopy.exe (Windows, x86)")):
    s, lv, ch, vd = summary(fn)
    rows.append([title, s, lv, ch, vd])
TABLE("Результаты ИАС для разных файлов", ["Файл", "Балл", "Уровень", "Цепочек", "Вывод ИАС"],
      rows, [4.2, 1.4, 2.2, 1.7, 7.0], size=10)
P("ИАС уверенно отделяет образец ВПО от легитимных программ. Единственный файл с вердиктом «вредоносная "
  "программа деструктивного типа» — Whiter.a. Демонстрационный образец sample_dropper, содержащий сигнатуры "
  "внедрения кода, получает высокий балл и вывод о необходимости углублённого анализа. Приложение .NET "
  "распознаётся как стандартное, как и в ПР №7.")
P("Наибольший балл среди легитимных файлов у cmd.exe (53). Это ожидаемо: интерпретатор команд действительно "
  "умеет удалять, копировать и перемещать файлы, снимать атрибуты, перебирать диски — ИАС нашла у него "
  "рекурсивный обход каталогов, снятие атрибутов и перебор дисков. Однако цепочки «массовое уничтожение» "
  "у cmd.exe нет: операции перечисления и удаления находятся в разных функциях и связаны командами "
  "пользователя, а не жёстко в коде. Именно это отличает инструмент двойного назначения от вайпера.")
P("Проверка выявила и исправленные в ходе работы недостатки: двойной учёт однотипных цепочек из разных "
  "функций (у cmd.exe балл был завышен до 100), срабатывание правила decrypt на функцию DecryptFileW, "
  "перенаправление WOW64 в 32-битной сборке, неверное чтение коротких строк UTF-16.")
P("Ограничения ИАС:", keep=True)
L(["для 64-битного кода часть аргументов остаётся неизвестной: компилятор держит значения в регистрах, "
   "заполняемых далеко от вызова, и сохраняет их в стек",
   "упакованные и зашифрованные файлы анализируются только по заголовкам и импорту — код нужно сначала распаковать",
   "для сборок .NET анализируется только машинный код загрузчика, а логика программы на языке IL — нет",
   "границы функций определяются эвристически, а код разбирается линейно, без построения полного графа "
   "управления; запутанный (обфусцированный) код может нарушить разбор",
   "выводы статического анализа о поведении в конкретной среде (права, WOW64) следует подтверждать "
   "динамическим анализом в изолированной виртуальной машине"])

# ================================================================ заключение

H1("ЗАКЛЮЧЕНИЕ")
P("В курсовой работе исследовано взаимодействие вредоносной программы с файловой системой на примере образца "
  "Whiter.a, и для этого разработана информационно-аналитическая система безопасности.")
P("ИАС создана по техническому заданию практической работы №8 на основе средства IncidentMap из ПР №7. Она "
  "принимает путь к исполняемому файлу или его содержимое через стандартный ввод, извлекает и нормализует "
  "строки, классифицирует их по матрице из 55 правил с привязкой к CWE, CVE и MITRE ATT&CK, оценивает риск "
  "и формирует отчёт из шести блоков. Главное отличие от предыдущих работ — модуль анализа взаимодействия "
  "с файловой системой: собственный декодер инструкций x86/x64 и модель стека и регистров позволяют "
  "восстановить аргументы вызовов WinAPI прямо из машинного кода и выявить 13 видов цепочек поведения.")
P("Анализ образца показал, что Whiter.a — деструктивный троян (вайпер) 2001 года, маскирующийся под "
  "генератор ключей Windows XP. При первом запуске он копирует себя в %SystemRoot%\\system32\\whismng.exe, "
  "прописывает автозапуск в HKLM\\...\\Run и показывает поддельный ключ в Блокноте. При следующем входе "
  "в систему он создаёт файл-шаблон c:\\wxp с текстом «You did a piracy, you deserve it.» и в бесконечном "
  "цикле обходит «Мои документы», все доступные для записи диски и диск C:, снимая с файлов защитные атрибуты, "
  "затирая их содержимое копией шаблона и удаляя. Все эти выводы получены без запуска образца и подтверждены "
  "конкретными адресами и аргументами вызовов.")
P("Проверка на легитимных программах показала, что ИАС отличает вредоносный код от инструментов двойного "
  "назначения: вывод о деструктивной программе сделан только для Whiter.a, а cmd.exe, notepad.exe и xcopy.exe "
  "получили объяснимые оценки. Сформулированы индикаторы компрометации и рекомендации по реагированию и "
  "защите. Поставленные задачи выполнены, цель работы достигнута.")
P("Развитие работы возможно в направлении поддержки распаковки популярных упаковщиков, построения полного "
  "графа управления функций и автоматической генерации правил YARA по найденным индикаторам.")

# ================================================================ источники

H1("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
sources = [
    "ГОСТ Р 50922-2006. Защита информации. Основные термины и определения. — М.: Стандартинформ, 2008. — 12 с.",
    "Sikorski M., Honig A. Practical Malware Analysis: The Hands-On Guide to Dissecting Malicious Software. — "
    "San Francisco: No Starch Press, 2012. — 800 p.",
    "Yosifovich P., Ionescu A., Russinovich M., Solomon D. Windows Internals. Part 1. 7th ed. — Redmond: "
    "Microsoft Press, 2017. — 800 p.",
    "Eagle C. The IDA Pro Book. 2nd ed. — San Francisco: No Starch Press, 2011. — 672 p.",
    "MITRE ATT&CK. Enterprise Matrix [Электронный ресурс]. — URL: https://attack.mitre.org/ "
    "(дата обращения: 30.10.2026).",
    "Common Weakness Enumeration (CWE) [Электронный ресурс]. — URL: https://cwe.mitre.org/ "
    "(дата обращения: 30.10.2026).",
    "National Vulnerability Database (NVD) [Электронный ресурс]. — URL: https://nvd.nist.gov/ "
    "(дата обращения: 30.10.2026).",
    "Microsoft. PE Format [Электронный ресурс]. — URL: https://learn.microsoft.com/windows/win32/debug/pe-format "
    "(дата обращения: 30.10.2026).",
    "Microsoft. File Management Functions [Электронный ресурс]. — URL: "
    "https://learn.microsoft.com/windows/win32/fileio/file-management-functions (дата обращения: 30.10.2026).",
    "Microsoft. File System Redirector [Электронный ресурс]. — URL: "
    "https://learn.microsoft.com/windows/win32/winprog64/file-system-redirector (дата обращения: 30.10.2026).",
    "Intel 64 and IA-32 Architectures Software Developer's Manual. Volume 2: Instruction Set Reference "
    "[Электронный ресурс]. — URL: https://www.intel.com/content/www/us/en/developer/articles/technical/"
    "intel-sdm.html (дата обращения: 30.10.2026).",
    "Capstone — The Ultimate Disassembly Framework [Электронный ресурс]. — URL: https://www.capstone-engine.org/ "
    "(дата обращения: 30.10.2026).",
]
for i, s_ in enumerate(sources, 1):
    P(f"{i}. {s_}")

# ================================================================ приложения

def appendix(letter, title):
    h = doc.add_heading("", level=1)
    h.add_run(f"ПРИЛОЖЕНИЕ {letter}")
    h.add_run().add_break()
    h.add_run(title)


appendix("А", "Отчёт ИАС по образцу Whiter.a")
P("Полный текстовый отчёт report_Whiter.a.exe.txt, сформированный ИАС.", indent=False)
for line in report:
    doc.add_paragraph(style="CodeSmall").add_run(line if line else " ")

appendix("Б", "Листинг главного модуля main.cpp")
for line in open(os.path.join(PROJ, "ExeIAS", "main.cpp"), encoding="utf-8").read().splitlines():
    doc.add_paragraph(style="CodeSmall").add_run(line.replace("\t", "    ") if line else " ")

appendix("В", "Листинг модуля анализа взаимодействия с файловой системой")
for fn in ("fsanalysis.h", "fsanalysis.cpp"):
    P(f"Файл {fn}", indent=False, bold=True)
    for line in open(os.path.join(PROJ, "ExeIAS", fn), encoding="utf-8").read().splitlines():
        doc.add_paragraph(style="CodeSmall").add_run(line.replace("\t", "    ") if line else " ")

doc.save(OUT)
print("saved", OUT, "figures", fig_no[0], "tables", tab_no[0], "listings", lst_no[0])
